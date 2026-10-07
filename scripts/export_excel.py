# -*- coding: utf-8 -*-
"""Readable Excel delivery for audited Baymax plans; source calculations stay unchanged."""
import argparse
import csv
import json
import math
from datetime import date, datetime, timedelta
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

NUTRIENTS = [('kcal','热量（kcal）'),('protein','蛋白质（g）'),('carbs','碳水（g）'),('fat','脂肪（g）'),('fiber','纤维（g）'),('calcium','钙（mg）'),('iron','铁（mg）'),('sodium','钠（mg）')]
RECORD_FIELDS = {'date':'日期','morning_weight_kg':'晨重（kg）','waist_cm':'腰围（cm）','actual_meals':'实际吃了什么','substitutions':'替换或外食','actual_kcal':'实际热量（kcal）','intake_complete':'全天饮食已完整记录','hunger':'饥饿感','energy':'精力','sleep_hours':'睡眠（小时）','symptoms':'不适或备注','actual_spend':'实际花费（元）'}
LABELS = {'age':'年龄','sex':'性别','height_cm':'身高（cm）','weight_kg':'当前体重（kg）','goal':'主要目标','target_weight_kg':'长期目标体重（kg）','target_body_fat_percent':'目标体脂（%）','health':'健康情况','work_activity':'工作活动','exercise_sessions_week':'规律活动次数/周','exercise_minutes_session':'每次活动分钟','wake_time':'起床时间','meals':'用餐规律','country':'国家','city':'城市','servings':'人数','daily_logging':'每日记录'}
PRICE_STATUS = {'verified':'已核实报价','estimated':'预算估价，非实时价','unknown':'价格未知'}


def read_records(path):
    """Normalize existing actual records; never fill from the planned menu."""
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        with path.open(encoding='utf-8-sig',newline='') as f:
            return list(csv.DictReader(f))
    wb = load_workbook(path,read_only=True,data_only=False)
    try:
        if '每日记录' not in wb.sheetnames:
            raise ValueError('Excel workbook missing 每日记录 sheet')
        rows = wb['每日记录'].iter_rows(min_row=4,values_only=True)
        header = next(rows)
        indexes = {}
        for key, label in RECORD_FIELDS.items():
            if header.count(label) != 1:
                raise ValueError('Excel journal missing or duplicate header: '+label)
            indexes[key] = header.index(label)
        result=[]
        for cells in rows:
            if all(c is None for c in cells):
                continue
            row={}
            for key,index in indexes.items():
                v=cells[index]
                if isinstance(v,(date,datetime)): v=v.isoformat()[:10]
                row[key]='' if v is None else str(v)
            if row['intake_complete'] in ('是','True'):row['intake_complete']='true'
            if row['intake_complete'] in ('否','False'):row['intake_complete']='false'
            result.append(row)
        return result
    finally:
        wb.close()


def plain(value):
    if value is None: return '未知'
    if isinstance(value,bool): return '是' if value else '否'
    if isinstance(value,(dict,list)): return json.dumps(value,ensure_ascii=False)
    return value


def textcell(cell, value):
    cell.value = value
    if isinstance(value,str): cell.data_type='s'  # Do not execute text that begins with '='.


def table(wb,name,headers,rows,widths,subtitle,tab='256B61',height_columns=None):
    ws=wb.create_sheet(name);ws.sheet_properties.tabColor=tab
    ws.sheet_view.zoomScale=80;ws.sheet_view.showGridLines=False
    count=len(headers)
    for row,title,color in [(1,'Baymax｜'+name,'183F39'),(2,subtitle,'EDF5F2')]:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=count)
        c=ws.cell(row,1);textcell(c,title);c.fill=PatternFill('solid',fgColor=color)
        c.font=Font(name='微软雅黑',size=18 if row==1 else 11,bold=row==1,color='FFFFFF' if row==1 else '35564F')
        c.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=34;ws.row_dimensions[2].height=38
    for col,header in enumerate(headers,1):
        c=ws.cell(4,col,header);c.fill=PatternFill('solid',fgColor='256B61');c.font=Font(name='微软雅黑',bold=True,color='FFFFFF',size=11)
        c.alignment=Alignment(vertical='center',wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width=widths[col-1]
    ws.row_dimensions[4].height=30
    for rownum,row in enumerate(rows,5):
        height=1
        for col,value in enumerate(row,1):
            c=ws.cell(rownum,col);textcell(c,value)
            c.font=Font(name='微软雅黑',size=11,color='243C37')
            c.fill=PatternFill('solid',fgColor='F3F7F5' if rownum%2 else 'FFFFFF')
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(value,(float,int)) and not isinstance(value,bool):c.number_format='0.0'
            if isinstance(value,str) and (height_columns is None or col<=height_columns):
                # Chinese glyphs are wider than ASCII; estimate wrapped line count.
                linecount=sum(max(1,math.ceil(sum(2 if ord(ch)>255 else 1 for ch in line)/max(8,widths[col-1]-2))) for line in value.split('\n'))
                height=max(height,linecount)
        ws.row_dimensions[rownum].height=min(360,max(30,height*16+10))
    ws.freeze_panes='A5';ws.auto_filter.ref='A4:'+get_column_letter(count)+str(max(4,4+len(rows)))
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:4';ws.print_options.horizontalCentered=True
    ws.print_area='A1:'+get_column_letter(count)+str(max(4,4+len(rows)))
    ws.oddFooter.center.text='Baymax · 第 &P 页 / 共 &N 页'
    return ws


def amount(g):
    return ('%.2fkg'%(g/1000)).rstrip('0').rstrip('.')+'（采购重量）' if g>=1000 else '%g克'%g


def ingredients(items):
    return '\n'.join('%s %g克'%(i['food'],i['grams']) for i in items)


def export_workbook(result,output,journal=None):
    output=Path(output)
    if output.exists():raise ValueError('refuse overwriting an existing Excel workbook or its user records')
    days=result['days'];wb=Workbook();wb.remove(wb.active)
    start=date.fromisoformat(days[0]['date'])
    means={k:round(sum(d['nutrients'][k]['value'] for d in days)/len(days),1) if all(d['nutrients'][k]['value'] is not None for d in days) else '未知' for k,_ in NUTRIENTS}
    overview=[['欢迎','我是大白。先看“每天吃什么”；采购时看“买菜清单”；周末照“周末备餐”操作；每天只填“每日记录”。'],
              ['计划日期',days[0]['date']+' 至 '+days[-1]['date']],['计划状态',result['status']],
              ['每日平均热量（估算）',str(means['kcal'])+' kcal'],['每日平均蛋白质（估算）',str(means['protein'])+' g'],
              ['预计采购预算',('未知' if result['estimated_purchase_total'] is None else '约 %.0f 元'%result['estimated_purchase_total'])+'；估价和历史均价不是门店实价。'],
              ['数字怎么理解','所有营养均为食物成分估算。未知不是0；核算通过不保证体重/体脂结果。未匹配商品和未核实价格在“食材依据”标明。'],
              ['重量怎么称','米、糙米、燕麦和干豆按干重；其余按各自可食生重或净重。不要把熟饭克数当干米克数；奶的ml和g不能直接等同。'],
              ['保存与设备','按逐餐说明、备餐步骤和原食品标签处理。设备、冷冻容量及食物限制未确认时仍是草案；不把一周熟饭冷藏。'],
              ['需要填写的地方','“买菜清单”最后一列可标记已买；“每日记录”浅蓝色格子填真实摄入和感受。留空表示没记录，不能把计划当成实际吃过。']]
    profile=result.get('profile') or {}
    if isinstance(profile,dict):
        for k,v in profile.items():overview.append([LABELS.get(k,k),plain(v)])
    overview += [['实施假设',str(a)] for a in result.get('assumptions',[])]
    for batch in result['operational_plan']:
        costs=[s['cost'] for s in result['shopping'] if s['purchase_date']==batch['purchase_date']]+[batch.get('delivery_cost')]
        subtotal='未知' if any(c is None for c in costs) else '约 %.0f 元'%sum(costs)
        overview.append([batch['purchase_date']+'采购预算',subtotal+'；覆盖'+batch['start']+'—'+batch['end']+'，非门店实价'])
    target_labels={'kcal':'试行热量（kcal）','protein_g':'蛋白质目标（g）','fiber_g':'纤维目标（g）','carbs_g':'碳水目标（g）','fat_g':'脂肪目标（g）','bmr_estimated_kcal':'基础代谢估算（kcal）','tdee_estimated_kcal':'每日消耗估算（kcal）','basis':'目标估算依据'}
    for k,v in (result.get('target') or {}).items():
        overview.append([target_labels.get(k,k), '—'.join(str(x) for x in v) if isinstance(v,list) else plain(v)])
    links=[['去看：'+s,'点击这一格打开相应工作表'] for s in ['每天吃什么','买菜清单','周末备餐','每日记录','阶段复盘','每日营养','分装标签','食材依据']]
    overview=overview[:3]+links+overview[3:]
    ws=table(wb,'开始使用',['项目','说明'],overview,[27,110],'一个文件就够了：吃什么 → 买什么 → 怎么准备 → 每天记录')
    for row in range(5,ws.max_row+1):
        label=ws.cell(row,1).value
        if isinstance(label,str) and label.startswith('去看：'):
            ws.cell(row,2).hyperlink="#'%s'!A1"%label[3:];ws.cell(row,2).font=Font(name='微软雅黑',size=11,color='187064',underline='single')
    menu=[]
    for d in days:
        for m in d['meals']:
            menu.append([d['date'],m['type'],m.get('time',''),m['name'],ingredients(m['ingredients']),
                         '\n'.join('%d. %s'%(n+1,st) for n,st in enumerate(m['steps'])),m.get('minutes','未填写'),
                         plain(m['nutrients']['kcal']['value']),plain(m['nutrients']['protein']['value']),m.get('storage',''),m.get('batch','当天现做')])
    ws=table(wb,'每天吃什么',['日期','餐次','时间','吃什么','食材与克数','怎么做','分钟','热量（kcal）','蛋白质（g）','保存/取用说明','分装标签'],menu,[13,8,8,25,43,50,8,12,12,70,35],'每天三餐完整列出；食材克数含油盐。展开上方“＋”可看估算营养和保存细节。',height_columns=7)
    ws.column_dimensions.group('H','K',hidden=True)
    shopping=[]
    for s in result['shopping']:
        shopping.append([s['purchase_date'],s['start']+'—'+s['end'],s['food'],
                         '%d份 × %g克 = %s'%(s['packages'],s['pack_g'],amount(s['purchased_g'])),
                         plain(s['cost']),PRICE_STATUS.get(s['price_status'],s['price_status']),s.get('channel') or '渠道未核实',None,
                         round(s['edible_use_g'],1),round(s['stock_before_g'],1),round(s['stock_after_g'],1),s.get('storage') or '待确认',s.get('evidence') or '没有实时商品证据',s.get('conditions') or '条件未核实'])
    ws=table(wb,'买菜清单',['采购日','供哪几天吃','食材','建议购买量','预算（元）','价格状态','建议渠道','是否已买','菜单净用量（g）','此前库存（g）','预计剩余（g）','保存方法','价格证据','购买条件'],shopping,[13,24,25,35,12,24,40,12,18,18,18,70,60,50],'“建议购买量”按包装/称重粒度估算；实物标签及收据优先，展开上方“＋”可看库存和保存说明。',height_columns=8)
    ws.column_dimensions.group('I','N',hidden=True)
    validation=DataValidation(type='list',formula1='"未买,已买,需替换"',allow_blank=True);ws.add_data_validation(validation);validation.add('H5:H'+str(ws.max_row))
    for row in range(5,ws.max_row+1):ws.cell(row,8).fill=PatternFill('solid',fgColor='EAF4FF')
    prep=[]
    for b in result['operational_plan']:
        for step in b['prep']:
            if isinstance(step,str):prep.append([b['purchase_date'],b['start']+'—'+b['end'],'','',step])
            else:prep.append([b['purchase_date'],b['start']+'—'+b['end'],step.get('time',''),step.get('title',''),step.get('detail','')])
    table(wb,'周末备餐',['采购/准备日','覆盖日期','什么时候','做什么','具体怎么做'],prep,[14,25,22,25,110],'按顺序完成。先确认冰箱和冷冻容量；米饭每日现煮，后半周易腐原料及时冷冻。')
    nutrition=[]
    for d in days:
        notes=['%s未知（已知小计%.1f，缺%d项）'%(label,d['nutrients'][k]['known_subtotal'],d['nutrients'][k]['missing_items']) for k,label in NUTRIENTS if d['nutrients'][k]['value'] is None]
        nutrition.append([d['date']]+[plain(d['nutrients'][k]['value']) for k,_ in NUTRIENTS]+[plain(d.get('estimated_consumed_cost')),'；'.join(notes) or '全部为数据库估算；实际商品仍需匹配'])
    table(wb,'每日营养',['日期']+[label for _,label in NUTRIENTS]+['食材消耗成本（元）','数据说明'],nutrition,[13]+[15]*8+[22,55],'这是计划营养，不是实际摄入。消耗成本与整包装采购预算分开；未知保留“未知”。')
    actual={}
    if journal:
        for row in read_records(journal):
            d=row.get('date')
            if d in actual:raise ValueError('duplicate actual record date')
            if d not in {x['date'] for x in days}:raise ValueError('record outside plan dates; preserve separately before changing periods')
            actual[d]=row
    records=[]
    for d in days:
        row=actual.get(d['date'],{})
        values=[d['date']]
        for key in list(RECORD_FIELDS)[1:]:
            value=row.get(key) or None
            if key=='intake_complete' and value: value='是' if value.lower()=='true' else ('否' if value.lower()=='false' else value)
            if key in ('morning_weight_kg','waist_cm','actual_kcal','sleep_hours','actual_spend') and value is not None:
                value=float(value)
                if not math.isfinite(value) or value<0:raise ValueError('invalid actual record number')
            values.append(value)
        records.append(values)
    ws=table(wb,'每日记录',list(RECORD_FIELDS.values()),records,[13,14,14,45,30,20,25,18,18,18,35,20],'只填真实情况。晨重可留空、腰围每周一次；“全天饮食已完整记录”不能由菜单完成自动推断。',tab='367BB1')
    for row in ws.iter_rows(min_row=5,min_col=2):
        for c in row:c.fill=PatternFill('solid',fgColor='EAF4FF');c.font=Font(name='微软雅黑',size=11,color='1D5D90')
    choice=DataValidation(type='list',formula1='"是,否"',allow_blank=True);ws.add_data_validation(choice);choice.add('G5:G34')
    reviews=[]
    for n in (7,14,21,30):
        reviews.append([n,str(start+timedelta(days=n-1)),
                        '实际吃了什么、饥饿精力、晨重趋势、腰围与采购花费',
                        '等待真实记录；体重差不是脂肪差。至少两周趋势后才讨论小幅热量调整。' if n!=30 else '等待真实30天记录；之后制定未来8–12周饮食规划，不编造效果。'])
    table(wb,'阶段复盘',['第几天','复盘日期','届时看什么','当前状态与下一步'],reviews,[12,15,60,90],'本表是复盘安排，不是已完成的监测，也不代表后台提醒已经启用。')
    portions=[]
    grouped={}
    for b in result.get('batch_rows',[]):grouped.setdefault(b['batch_id'],[]).append(b)
    for bid,items in grouped.items():
        b=items[0];portions.append([b['prepare'],b['eat'],bid,ingredients([{'food':i['food'],'grams':i['grams']} for i in items]),b['thaw'],b['storage']])
    table(wb,'分装标签',['准备日','食用日','标签','这一餐准备哪些食材','什么时候解冻','保存与烹饪'],portions,[13,13,30,48,65,90],'每个标签对应一人份；肉菜分开、米和调料干存、奶按标签保管，不把整餐混成一袋冷冻。')
    substitutions=[]
    for s in result.get('substitutions',[]):
        def totals(items):
            return {k:sum(result['sources'][i['food']]['food']['nutrients'][k]*i['grams']/100 for i in items) if all(result['sources'][i['food']]['food']['nutrients'][k] is not None for i in items) else None for k,_ in NUTRIENTS}
        a,b=totals(s['original']),totals(s['replacement'])
        def budget(items):
            total=0
            for i in items:
                spec=result['sources'][i['food']]['spec'];offer=spec['offer']
                if offer.get('pack_price') is None:return None
                total+=i['grams']/spec['edible_yield']/offer['pack_g']*offer['pack_price']
            return round(total,2)
        substitutions.append([ingredients(s['original']),ingredients(s['replacement']),plain(a['kcal']),plain(b['kcal']),plain(a['protein']),plain(b['protein']),plain(budget(s['original'])),plain(budget(s['replacement'])),s.get('note','整日重算，非全面营养等价')])
    if substitutions:table(wb,'临时替换',['原食材','可以换成','原热量（kcal）','替换热量（kcal）','原蛋白质（g）','替换蛋白质（g）','原预算（元）','替换预算（元）','注意事项'],substitutions,[40,40,18,20,20,20,18,18,70],'只有无相关食物限制时才替换；来源与价格未核实，替换后整日营养需重新核算。')
    sources=[]
    for alias,e in result['sources'].items():
        f=e['food'];s=e['spec'];src=f['source'];offer=s['offer']
        sources.append([alias,'参考，实际食物未匹配' if s['match'] in ('reference','unresolved') else ('已核对商品标签' if s['match']=='exact-label' else '通用食物匹配估算'),f['description'],s['weight_basis'],
                        *[plain(f['nutrients'][k]) for k,_ in NUTRIENTS],src.get('publisher') or src.get('brand','未知'),src.get('version','标签版本见ID'),src.get('accessed','未知'),f['id'],src.get('url') or src.get('evidence','未知'),s.get('match_evidence','未知'),PRICE_STATUS.get(offer['status'],offer['status']),offer.get('price_basis') or offer.get('evidence') or '没有核实报价证据'])
    table(wb,'食材依据',['食材','匹配状态','数据库原名称','重量口径']+[label+' /100g' for _,label in NUTRIENTS]+['来源机构','来源版本','访问日','条目ID','来源地址或证据','匹配说明','价格状态','价格依据'],sources,[28,28,65,40]+[20]*8+[20,25,15,80,80,65,28,70],'需要核对数据时再看这里；所有来源和未知值保留。哈希和完整审计仍保存在本机JSON底稿。',tab='87958F')
    gaps=[[g] for g in result['gaps']]
    table(wb,'待核实事项',['尚未解决的数据或执行条件'],gaps,[120],'这些缺口没有因为导出Excel而消失；不能把文件格式改进当作数据精度提高。',tab='C99B44')
    wb.active=0;output.parent.mkdir(parents=True,exist_ok=True);wb.save(output);wb.close()
    return output


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--audit',required=True);p.add_argument('--out',required=True);p.add_argument('--journal')
    a=p.parse_args();result=json.loads(Path(a.audit).read_text(encoding='utf-8'))
    print(export_workbook(result,a.out,a.journal))

if __name__=='__main__':main()
