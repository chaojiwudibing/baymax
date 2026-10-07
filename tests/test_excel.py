# -*- coding: utf-8 -*-
import copy
import json
import sys
import unittest
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import test_nutrition
from calculate_plan import calculate, deliver
from export_excel import export_workbook, read_records
from review_logs import review

class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.fixture=test_nutrition.NutritionTests();self.fixture.setUp()
        self.root=self.fixture.root;self.plan=self.fixture.plan;self.db=self.fixture.db
    def tearDown(self):self.fixture.tearDown()
    def result(self):return calculate(self.plan,self.db)
    def test_default_delivery_real_workbook_and_readable_layout(self):
        deliver(self.result(),self.root/'out');p=self.root/'out/Baymax-30天饮食计划.xlsx'
        self.assertTrue(p.exists());w=load_workbook(p)
        self.assertEqual(w.active.title,'开始使用');self.assertEqual(w['每天吃什么'].max_row,94)
        self.assertEqual(w['每日营养'].max_row,34);self.assertEqual(w['每日记录'].max_row,34)
        self.assertEqual(w['每天吃什么']['E4'].value,'食材与克数')
        self.assertEqual(w['每日营养']['B5'].value,300)
        self.assertEqual(w['每日记录']['B5'].value,None)
        self.assertEqual(w['每天吃什么'].freeze_panes,'A5');self.assertTrue(w['每天吃什么']['E5'].alignment.wrap_text)
        self.assertTrue(any(c.hyperlink for row in w['开始使用'] for c in row));w.close()
    def test_missing_nutrient_and_price_remain_unknown(self):
        import sqlite3
        c=sqlite3.connect(self.db);c.execute("delete from nutrients where name='calcium'");c.commit();c.close()
        self.plan['foods']['test']['offer']['pack_price']=None
        p=self.root/'out.xlsx';export_workbook(self.result(),p);w=load_workbook(p)
        self.assertEqual(w['每日营养']['G5'].value,'未知')
        self.assertIn('钙',w['每日营养']['K5'].value);self.assertEqual(w['买菜清单']['E5'].value,'未知')
        w.close()
    def test_untrusted_text_is_not_an_excel_formula(self):
        self.plan['days'][0]['meals'][0]['name']='=HYPERLINK("bad")'
        p=self.root/'out.xlsx';export_workbook(self.result(),p);w=load_workbook(p,data_only=False)
        self.assertEqual(w['每天吃什么']['D5'].data_type,'s');self.assertEqual(w['每天吃什么']['D5'].value,'=HYPERLINK("bad")');w.close()
    def test_overwrite_rejected(self):
        p=self.root/'out.xlsx';export_workbook(self.result(),p)
        with self.assertRaises(ValueError):export_workbook(self.result(),p)
    def test_actual_records_survive_revision_and_support_review(self):
        old=self.root/'old.xlsx';export_workbook(self.result(),old);w=load_workbook(old)
        ws=w['每日记录'];ws['B5']=74;ws['D5']='实际早餐';ws['F5']=2500;ws['G5']='是';w.save(old);w.close()
        new=self.root/'new.xlsx';export_workbook(self.result(),new,journal=old)
        rows=read_records(new);self.assertEqual(rows[0]['morning_weight_kg'],'74');self.assertEqual(rows[0]['actual_meals'],'实际早餐')
        self.assertEqual(rows[1]['morning_weight_kg'],'')
        stats=review(new,'2020-01-01');self.assertEqual(stats['complete_intake_days'],1)
        self.assertEqual(stats['baseline_first7']['sample_count'],1);self.assertEqual(stats['mean_reported_complete_intake_kcal'],2500)
    def test_duplicate_record_and_missing_header_rejected(self):
        p=self.root/'old.xlsx';export_workbook(self.result(),p);w=load_workbook(p);w['每日记录']['A6']='2020-01-01';w.save(p);w.close()
        with self.assertRaises(ValueError):export_workbook(self.result(),self.root/'new.xlsx',journal=p)
        w=load_workbook(p);w['每日记录']['B4']='错误标题';w.save(p);w.close()
        with self.assertRaises(ValueError):read_records(p)
    def test_substitution_is_recomputed_from_source(self):
        self.plan['substitutions']=[{'original':[{'food':'test','grams':100}],'replacement':[{'food':'test','grams':200}],'note':'synthetic'}]
        p=self.root/'out.xlsx';export_workbook(self.result(),p);w=load_workbook(p)
        self.assertEqual(w['临时替换']['C5'].value,100);self.assertEqual(w['临时替换']['D5'].value,200);w.close()
    def test_macro_source_unknown_still_blocks_strict_status(self):
        self.plan['foods']['test']['match']='reference';r=self.result();p=self.root/'out.xlsx';export_workbook(r,p);w=load_workbook(p)
        values=[c.value for row in w['开始使用'] for c in row]
        self.assertIn('草案：数据缺口未解决',values)
        self.assertEqual(w['待核实事项']['A5'].value,r['gaps'][0]);w.close()

if __name__=='__main__':unittest.main()
